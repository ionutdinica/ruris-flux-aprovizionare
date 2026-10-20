"""Importul comenzilor din Excel (cerința: „Achizițiile țin într-un Excel ce au comandat”).

Trei pași, fiecare cu responsabilitatea lui:
  1. read_rows(fișier)        – citește rândurile; tolerant la formate (date, numere cu virgulă);
  2. plan_import(rânduri, ...) – validează TOT și grupează liniile pe comenzi; erorile sunt
                                 raportate pe rând, pe înțelesul omului; nu salvează nimic;
  3. execute_import(plan, ...) – creează comenzile ca ciorne, toate sau niciuna (atomic).

Între pașii 2 și 3 utilizatorul vede o previzualizare. La confirmare, rândurile sunt
validate din nou: între timp un produs putea fi dezactivat sau comanda deja importată.

Regulile de business NU sunt rescrise aici: fiecare comandă și linie construită din
Excel trece prin aceleași validări ale modelului (full_clean) ca formularul din aplicație.
"""

import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from apps.audit.models import BusinessEvent
from apps.audit.services import record_event
from apps.catalog.models import Product, Supplier
from apps.core.models import Currency
from apps.core.text import fold
from apps.core.validators import normalize_cui

from .models import ZERO, PurchaseOrder, PurchaseOrderLine

MAX_ROWS = 1000  # linii de comandă într-un fișier
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024  # protecție la „arhive bombă” (.xlsx e o arhivă zip)
SHEET_NAME = "Comenzi"


class ImportFileError(Exception):
    """Fișierul nu poate fi folosit deloc (format greșit, coloane lipsă, prea mare)."""


@dataclass(frozen=True)
class ImportColumn:
    key: str
    title: str
    required: bool
    help: str
    example: str = ""


# Coloanele fișierului, în ordinea din șablon. Antetul e recunoscut după nume
# (fără diacritice, majuscule sau semne), deci ordinea coloanelor nu contează.
COLUMNS = (
    ImportColumn(
        "reference",
        "Referință comandă",
        True,
        "Rândurile cu aceeași referință formează o singură comandă "
        "(ex. numărul comenzii din Excel-ul vechi).",
        "EXC-0017",
    ),
    ImportColumn(
        "supplier",
        "Cod fiscal furnizor",
        True,
        "CUI-ul (cu sau fără RO) sau codul de TVA al furnizorului, ca în nomenclator "
        "(vezi foaia Furnizori).",
        "RO14399840",
    ),
    ImportColumn(
        "order_date",
        "Data comenzii",
        False,
        "Implicit: data importului. Format: zz.ll.aaaa.",
        "15.10.2026",
    ),
    ImportColumn("currency", "Monedă", False, "RON sau EUR. Implicit: moneda furnizorului.", "RON"),
    ImportColumn(
        "product",
        "Cod produs",
        True,
        "Codul din nomenclatorul de produse (vezi foaia Produse).",
        "MOT-65",
    ),
    ImportColumn("quantity", "Cantitate", True, "Număr pozitiv, maxim 3 zecimale.", "10"),
    ImportColumn(
        "unit_price", "Preț unitar", True, "Fără TVA, în moneda comenzii; maxim 4 zecimale.", "1450"
    ),
    ImportColumn(
        "vat_rate", "Cotă TVA (%)", False, "21, 11 sau 0. Implicit: cota produsului.", "21"
    ),
    ImportColumn(
        "expected_date",
        "Data estimată",
        True,
        "Când trebuie să sosească linia. Format: zz.ll.aaaa.",
        "05.11.2026",
    ),
    ImportColumn("notes", "Observații", False, "Se trec pe comandă (prima valoare completată)."),
)
TITLES = {column.key: column.title for column in COLUMNS}


# --- Pasul 1: citirea fișierului ------------------------------------------------------------


def _header_key(text):
    """„Cotă TVA (%)” → „cotatva”: antetul e recunoscut oricum ar fi scris."""
    return re.sub(r"[^a-z0-9]", "", fold(str(text or "")))


def _check_archive(uploaded_file):
    """Un .xlsx e o arhivă zip; refuzăm fișierele care s-ar „umfla” enorm la dezarhivare."""
    try:
        with zipfile.ZipFile(uploaded_file) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as error:
        raise ImportFileError("Fișierul nu este un Excel .xlsx valid.") from error
    finally:
        uploaded_file.seek(0)
    if total > MAX_UNCOMPRESSED_BYTES:
        raise ImportFileError("Fișierul este prea mare.")


def _plain(value):
    """Valoarea unei celule, într-o formă care poate fi păstrată în sesiune (JSON)."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | float):
        return value  # numerele rămân numere (vezi parse_decimal)
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def read_rows(uploaded_file):
    """Citește rândurile cu date: [{"row": 2, "values": {"product": "MOT-65", ...}}, ...]."""
    _check_archive(uploaded_file)
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception as error:  # orice fișier corupt sau în alt format
        raise ImportFileError(
            "Fișierul nu poate fi citit. Folosește un fișier Excel .xlsx (de exemplu șablonul)."
        ) from error

    try:
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
        sheet.reset_dimensions()  # unele programe salvează dimensiuni greșite ale foii
        rows = sheet.iter_rows(values_only=True)
        positions = _column_positions(next(rows, None))

        result = []
        for number, cells in enumerate(rows, start=2):
            values = {
                key: _plain(cells[i]) if i < len(cells) else None for key, i in positions.items()
            }
            if all(value is None for value in values.values()):
                continue  # rând gol
            if len(result) >= MAX_ROWS:
                raise ImportFileError(f"Fișierul are mai mult de {MAX_ROWS} de linii.")
            result.append({"row": number, "values": values})
    finally:
        workbook.close()

    if not result:
        raise ImportFileError("Fișierul nu conține nicio linie de comandă.")
    return result


def _column_positions(header):
    found = {_header_key(title): index for index, title in enumerate(header or ()) if title}
    positions = {
        column.key: found[_header_key(column.title)]
        for column in COLUMNS
        if _header_key(column.title) in found
    }
    missing = [
        column.title for column in COLUMNS if column.required and column.key not in positions
    ]
    if missing:
        raise ImportFileError(
            f"Lipsesc coloanele: {', '.join(missing)}. Pe primul rând trebuie să fie antetul "
            "din șablon."
        )
    return positions


# --- Conversii tolerante -----------------------------------------------------------------------

THOUSANDS_WITH_DOTS = re.compile(r"^-?\d{1,3}(\.\d{3})+$")  # 1.500 sau 12.345.678


def parse_decimal(value):
    """Număr din celulă: numeric sau text în stil românesc ori englezesc.

    12,5 → 12.5 · 1.234,56 → 1234.56 · 1,234.56 → 1234.56 · 1.500 → 1500 (mii)
    Celulele numerice sunt preferate: textul „1.500” e interpretat ca o mie cinci sute.
    """
    if value is None:
        return None
    if isinstance(value, int | float):
        # Excel păstrează numerele ca float: 0,1 + 0,2 poate apărea ca 0.30000000000000004.
        try:
            return Decimal(repr(value)).quantize(Decimal("0.000001")).normalize()
        except InvalidOperation as error:  # ex. 1e30: prea multe cifre
            raise ValueError(f"„{value}” este un număr prea mare.") from error

    text = str(value).replace(" ", "").replace(" ", "").removesuffix("%")
    if "," in text and "." in text:
        # Ambii separatori: cel din dreapta e cel zecimal.
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif text.count(",") == 1:
        text = text.replace(",", ".")  # virgulă zecimală
    elif "," in text:
        text = text.replace(",", "")  # 1,234,567
    elif THOUSANDS_WITH_DOTS.match(text):
        text = text.replace(".", "")  # 1.500 = o mie cinci sute
    try:
        return Decimal(text)
    except InvalidOperation as error:
        raise ValueError(f"„{value}” nu este un număr.") from error


DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y")


def parse_date(value):
    """Dată din celulă: dată Excel, număr de serie Excel sau text (21.10.2026, 2026-10-21...)."""
    if value is None:
        return None
    if isinstance(value, int | float) and 20000 < value < 100000:
        return from_excel(value).date()  # celulă de tip dată, dar fără format de dată
    text = str(value).strip()
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    raise ValueError(f"„{value}” nu este o dată validă (folosește zz.ll.aaaa).")


def parse_vat_rate(value):
    """21, „21%”, 0,21 (celulă formatată ca procent în Excel) → 21.00."""
    rate = parse_decimal(value)
    if rate is not None and 0 < rate < 1:
        rate *= 100
    return rate.quantize(Decimal("0.01")) if rate is not None else None


# --- Pasul 2: validare și plan --------------------------------------------------------------------


@dataclass
class PlannedOrder:
    reference: str
    first_row: int
    order: PurchaseOrder  # nesalvată
    lines: list = field(default_factory=list)  # [(rând, PurchaseOrderLine nesalvată)]

    @property
    def net_amount(self):
        return sum((line.net_amount for _, line in self.lines), ZERO)


@dataclass
class ImportPlan:
    orders: list = field(default_factory=list)
    errors: list = field(default_factory=list)  # [(rând, mesaj)]

    @property
    def is_valid(self):
        return bool(self.orders) and not self.errors

    @property
    def line_count(self):
        return sum(len(planned.lines) for planned in self.orders)

    def error_messages(self):
        return [f"Rândul {row}: {message}" for row, message in sorted(set(self.errors))]


def plan_import(rows, user):
    """Validează rândurile și le grupează pe comenzi. Nu salvează nimic."""
    plan = ImportPlan()
    suppliers = {s.tax_id: s for s in Supplier.objects.all()}
    products = {p.code: p for p in Product.objects.all()}
    today = timezone.localdate()

    parsed = []
    for raw in rows:
        values, errors = _parse_row(raw["values"], suppliers, products, today)
        plan.errors += [(raw["row"], message) for message in errors]
        parsed.append((raw["row"], values, not errors))

    groups = {}
    for row, values, valid in parsed:
        reference = values.get("reference") or ""
        groups.setdefault(reference, []).append((row, values, valid))

    for reference, group_rows in groups.items():
        if not reference:
            continue  # rândurile fără referință au primit deja eroare
        planned = _plan_order(reference, group_rows, user, plan.errors)
        if planned:
            plan.orders.append(planned)
    return plan


def _parse_row(values, suppliers, products, today):
    """Convertește valorile unui rând și verifică existența furnizorului și a produsului."""
    errors, result = [], {}

    def required(key):
        if values.get(key) is None:
            errors.append(f"{TITLES[key]}: lipsește.")
            return False
        return True

    def convert(key, parser):
        try:
            return parser(values.get(key))
        except ValueError as error:
            errors.append(f"{TITLES[key]}: {error}")
            return None

    if required("reference"):
        # Fără trunchiere: o referință prea lungă e raportată de validarea modelului.
        result["reference"] = str(values["reference"])

    if required("supplier"):
        supplier = suppliers.get(normalize_cui(str(values["supplier"])))
        if supplier is None:
            errors.append(
                f"furnizorul cu codul fiscal „{values['supplier']}” nu există în nomenclator."
            )
        elif not supplier.is_active:
            errors.append(f"furnizorul {supplier} este inactiv.")
        else:
            result["supplier"] = supplier

    if required("product"):
        product = products.get(str(values["product"]).strip().upper())
        if product is None:
            errors.append(f"produsul „{values['product']}” nu există în nomenclator.")
        elif not product.is_active:
            errors.append(f"produsul {product} este inactiv.")
        else:
            result["product"] = product

    currency = str(values.get("currency") or "").strip().upper()
    if currency and currency not in Currency.values:
        errors.append(f"{TITLES['currency']}: „{currency}” nu este acceptată (RON sau EUR).")
    elif currency:
        result["currency"] = currency
    elif "supplier" in result:
        result["currency"] = result["supplier"].default_currency

    result["order_date"] = convert("order_date", parse_date) or today
    if required("expected_date"):
        result["expected_date"] = convert("expected_date", parse_date)
    if required("quantity"):
        result["quantity"] = convert("quantity", parse_decimal)
    if required("unit_price"):
        result["unit_price"] = convert("unit_price", parse_decimal)
    result["vat_rate"] = convert("vat_rate", parse_vat_rate)
    if result["vat_rate"] is None and "product" in result:
        result["vat_rate"] = result["product"].vat_rate
    result["notes"] = str(values.get("notes") or "")
    return result, errors


# Câmpurile care trebuie să fie identice pe toate rândurile unei comenzi.
SHARED_FIELDS = ("supplier", "currency", "order_date")
# Câmpurile din care se construiește o linie de comandă.
LINE_FIELDS = ("product", "quantity", "unit_price", "vat_rate", "expected_date")


def _plan_order(reference, group_rows, user, errors):
    """Construiește comanda unei referințe și o validează cu regulile modelului.

    Raportează tot ce se poate verifica, chiar dacă unele rânduri au deja erori
    de citire: utilizatorul corectează fișierul o singură dată, nu eroare cu eroare.
    """
    group_ok = all(valid for _, _, valid in group_rows)

    # 1. Furnizorul, moneda și data trebuie să fie aceleași pe toate rândurile comenzii.
    #    Se compară doar valorile citite corect (o valoare greșită are deja eroarea ei).
    shared = {}
    for key in SHARED_FIELDS:
        rows_with_value = [(row, values[key]) for row, values, _ in group_rows if key in values]
        if not rows_with_value:
            continue
        first_row, first_value = rows_with_value[0]
        shared[key] = first_value
        for row, value in rows_with_value[1:]:
            if value != first_value:
                errors.append(
                    (
                        row,
                        f"{TITLES[key]} diferă de rândul {first_row}, deși referința "
                        f"„{reference}” este aceeași (o comandă are un singur furnizor, "
                        "o singură monedă și o singură dată).",
                    )
                )
                group_ok = False

    dates = [values["expected_date"] for _, values, _ in group_rows if values.get("expected_date")]
    order = PurchaseOrder(
        supplier=shared.get("supplier"),
        owner=user,
        currency=shared.get("currency", ""),
        order_date=shared["order_date"],
        expected_date=min(dates) if dates else None,
        external_reference=reference,
        notes=next((values["notes"] for _, values, _ in group_rows if values["notes"]), ""),
    )
    planned = PlannedOrder(reference=reference, first_row=group_rows[0][0], order=order)

    # 2. Fiecare linie trece prin aceleași reguli ca în formularul din aplicație
    #    (cantitate pozitivă, numărul de zecimale, cotă TVA permisă, termen după data comenzii).
    for row, values, _ in group_rows:
        if any(values.get(key) is None for key in LINE_FIELDS):
            continue  # rândul are deja o eroare de citire pentru aceste câmpuri
        line = PurchaseOrderLine(order=order, **{key: values[key] for key in LINE_FIELDS})
        if _model_errors(line, row, errors, exclude=["order"]):
            group_ok = False
        planned.lines.append((row, line))

    # 3. Regulile comenzii (ex. referința externă deja folosită la acest furnizor).
    if group_ok and not _model_errors(order, planned.first_row, errors, exclude=["number"]):
        return planned
    return None


FIELD_TITLES = {**TITLES, "number": "Număr", "external_reference": "Referință comandă"}


def _model_errors(instance, row, errors, exclude):
    try:
        instance.full_clean(exclude=exclude)
    except ValidationError as error:
        for field_name, messages in error.message_dict.items():
            title = FIELD_TITLES.get(field_name)
            for message in messages:
                errors.append((row, f"{title}: {message}" if title else message))
        return True
    return False


# --- Pasul 3: crearea comenzilor -----------------------------------------------------------------


@transaction.atomic
def execute_import(plan, filename):
    """Creează toate comenzile planului, ca ciorne. Toate sau niciuna."""
    if not plan.is_valid:
        raise ValueError("Planul de import conține erori.")
    created = []
    for planned in plan.orders:
        order = planned.order
        order.save()
        for _, line in planned.lines:
            line.order = order
            line.save()
        created.append(order)

    record_event(
        BusinessEvent.Type.ORDERS_IMPORTED,
        f"Import din Excel: {len(created)} comenzi ({plan.line_count} linii) din „{filename}”.",
        data={"fișier": filename, "comenzi": ", ".join(order.number for order in created)},
    )
    return created
