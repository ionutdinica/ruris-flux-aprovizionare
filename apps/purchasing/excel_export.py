"""Fișierele Excel generate pentru comenzi: exportul listei și șablonul de import."""

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.catalog.models import Product, Supplier
from apps.core.excel import (
    DATE,
    DATETIME,
    MONEY,
    PRICE,
    QUANTITY,
    Column,
    style_header,
    write_table,
)
from apps.core.models import Currency, VatRate

from .excel_import import COLUMNS, MAX_ROWS, SHEET_NAME

# --- Exportul comenzilor ---------------------------------------------------------------------

ORDER_COLUMNS = [
    Column("Număr"),
    Column("Referință externă"),
    Column("Data comenzii", DATE),
    Column("Furnizor"),
    Column("Cod fiscal furnizor"),
    Column("Status"),
    Column("Termen", DATE),
    Column("Întârziată"),
    Column("Monedă"),
    Column("Valoare fără TVA", MONEY),
    Column("TVA", MONEY),
    Column("Total", MONEY),
    Column("Responsabil"),
    Column("Trimisă la", DATETIME),
    Column("Observații", width=40),
]

LINE_COLUMNS = [
    Column("Număr comandă"),
    Column("Status comandă"),
    Column("Furnizor"),
    Column("Cod produs"),
    Column("Produs"),
    Column("UM"),
    Column("Cantitate comandată", QUANTITY),
    Column("Recepționat", QUANTITY),
    Column("Facturat", QUANTITY),
    Column("Rest de recepționat", QUANTITY),
    Column("Preț unitar", PRICE),
    Column("Cotă TVA (%)", QUANTITY),
    Column("Valoare fără TVA", MONEY),
    Column("Monedă"),
    Column("Termen", DATE),
]


def build_orders_export(orders, lines, *, user, filters):
    """Registru cu trei foi: Comenzi (una pe rând), Linii (cu progresul) și Informații."""
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Comenzi"
    write_table(
        sheet,
        ORDER_COLUMNS,
        (
            [
                order.number,
                order.external_reference,
                order.order_date,
                order.supplier.name,
                order.supplier.tax_id,
                order.get_status_display(),
                order.expected_date,
                "Da" if getattr(order, "is_late", False) else "Nu",
                order.currency,
                order.net_amount,
                order.vat_amount,
                order.gross_amount,
                str(order.owner),
                order.sent_at,
                order.notes,
            ]
            for order in orders
        ),
    )

    write_table(
        workbook.create_sheet("Linii"),
        LINE_COLUMNS,
        (
            [
                line.order.number,
                line.order.get_status_display(),
                line.order.supplier.name,
                line.product.code,
                line.product.name,
                line.product.unit,
                line.quantity,
                line.received_qty,
                line.invoiced_qty,
                max(line.quantity - line.received_qty, 0),
                line.unit_price,
                line.vat_rate,
                line.net_amount,
                line.order.currency,
                line.expected_date,
            ]
            for line in lines
        ),
    )

    info = workbook.create_sheet("Informații")
    info.append(["Export comenzi – Aprovizionare RURIS"])
    info["A1"].font = Font(bold=True, size=14)
    for label, value in [
        ("Generat la", timezone.localtime().strftime("%d.%m.%Y %H:%M")),
        ("Generat de", str(user)),
        ("Filtre aplicate", filters or "fără filtre (toate comenzile vizibile)"),
        ("Număr de comenzi", len(orders)),
    ]:
        info.append([label, value])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 60
    return workbook


# --- Șablonul de import ---------------------------------------------------------------------


def build_import_template():
    """Șablonul de import, generat din datele curente.

    Foaia „Comenzi” are antetul așteptat și liste derulante (furnizori, produse, monede,
    cote TVA), alimentate din foile „Furnizori” și „Produse”. Foaia „Instrucțiuni”
    explică fiecare coloană și conține un exemplu.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append([column.title for column in COLUMNS])
    style_header(sheet)
    sheet.freeze_panes = "A2"
    keys = [column.key for column in COLUMNS]

    def letter(key):
        return get_column_letter(keys.index(key) + 1)

    for column in COLUMNS:
        sheet.column_dimensions[letter(column.key)].width = max(len(column.title) + 4, 16)
    for key in ("order_date", "expected_date"):
        for row in range(2, MAX_ROWS + 2):
            sheet[f"{letter(key)}{row}"].number_format = DATE

    suppliers = list(Supplier.objects.active().order_by("name"))
    products = list(Product.objects.active().order_by("code"))
    last_row = MAX_ROWS + 1

    def dropdown(key, formula):
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.errorStyle = "warning"  # avertizează, dar permite și alte valori
        validation.error = "Valoarea nu este în listă. Verifică foaia de referință."
        sheet.add_data_validation(validation)
        validation.add(f"{letter(key)}2:{letter(key)}{last_row}")

    dropdown("supplier", f"=Furnizori!$B$2:$B${max(len(suppliers), 1) + 1}")
    dropdown("product", f"=Produse!$A$2:$A${max(len(products), 1) + 1}")
    dropdown("currency", f'"{",".join(Currency.values)}"')
    dropdown("vat_rate", f'"{",".join(str(int(float(v))) for v in VatRate.values)}"')

    _instructions_sheet(workbook.create_sheet("Instrucțiuni"))

    reference = workbook.create_sheet("Furnizori")
    write_table(
        reference,
        [Column("Denumire"), Column("Cod fiscal"), Column("Moneda implicită")],
        ([s.name, s.tax_id, s.default_currency] for s in suppliers),
    )
    reference = workbook.create_sheet("Produse")
    write_table(
        reference,
        [Column("Cod"), Column("Denumire"), Column("UM"), Column("Cotă TVA (%)", QUANTITY)],
        ([p.code, p.name, p.unit, p.vat_rate] for p in products),
    )
    return workbook


def _instructions_sheet(sheet):
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 80
    sheet.column_dimensions["D"].width = 16
    lines = [
        "Import comenzi din Excel",
        "",
        "1. Completează foaia „Comenzi”: un rând pentru fiecare linie de comandă.",
        "2. Rândurile cu aceeași „Referință comandă” formează o singură comandă "
        "(același furnizor, aceeași monedă, aceeași dată).",
        "3. Furnizorii și produsele trebuie să existe deja în aplicație (vezi foile "
        "„Furnizori” și „Produse”).",
        "4. Încarcă fișierul în aplicație: vei vedea o previzualizare și eventualele erori, "
        "pe rând, înainte de salvare.",
        "5. Comenzile se creează ca ciorne: verifică-le și trimite-le din aplicație.",
        "",
        "Coloane",
    ]
    for text in lines:
        sheet.append([text])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet[f"A{len(lines)}"].font = Font(bold=True)

    header_row = sheet.max_row + 1
    sheet.append(["Coloană", "Obligatorie", "Descriere", "Exemplu"])
    style_header(sheet, header_row)
    for column in COLUMNS:
        sheet.append([column.title, "da" if column.required else "nu", column.help, column.example])
        sheet.cell(row=sheet.max_row, column=3).alignment = Alignment(wrap_text=True)
