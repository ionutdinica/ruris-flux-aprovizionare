"""Exportul listei de comenzi în Excel."""

from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.audit.models import BusinessEvent
from apps.core.excel import XLSX_CONTENT_TYPE
from apps.purchasing.models import PurchaseOrder

from .factories import PurchaseOrderFactory, PurchaseOrderLineFactory

pytestmark = pytest.mark.django_db

EXPORT = reverse("purchasing:order_export")


def sent_order(**line_values):
    line = PurchaseOrderLineFactory(**line_values)
    PurchaseOrder.objects.filter(pk=line.order_id).update(status=PurchaseOrder.Status.SENT)
    return line.order


def download(client, **params):
    response = client.get(EXPORT, params)
    assert response["Content-Type"] == XLSX_CONTENT_TYPE
    assert response["Content-Disposition"].startswith('attachment; filename="comenzi_')
    return load_workbook(BytesIO(response.content))


def test_export_contains_orders_lines_and_info(client, procurement_user):
    order = sent_order(quantity=Decimal("3"), unit_price=Decimal("33.3333"))
    client.force_login(procurement_user)

    workbook = download(client)

    assert workbook.sheetnames == ["Comenzi", "Linii", "Informații"]
    header, values = list(workbook["Comenzi"].iter_rows(values_only=True))
    data = dict(zip(header, values, strict=True))
    assert data["Număr"] == order.number
    assert data["Valoare fără TVA"] == 100  # 3 × 33,3333 = 99,9999 → 100,00 – număr, nu text
    assert data["Status"] == "Trimisă"
    line_header, line_values = list(workbook["Linii"].iter_rows(values_only=True))
    assert dict(zip(line_header, line_values, strict=True))["Rest de recepționat"] == 3


def test_export_respects_filters(client, procurement_user):
    wanted = sent_order()
    PurchaseOrderFactory()  # ciornă
    client.force_login(procurement_user)

    workbook = download(client, status="SENT")

    numbers = [r[0] for r in workbook["Comenzi"].iter_rows(min_row=2, values_only=True)]
    assert numbers == [wanted.number]
    info = dict(workbook["Informații"].iter_rows(min_row=2, values_only=True))
    assert info["Filtre aplicate"] == "status: Trimisă"


def test_export_respects_visibility(client, accounting_user):
    """Contabilitatea nu vede ciornele nici în export."""
    visible = sent_order()
    PurchaseOrderFactory()
    client.force_login(accounting_user)

    workbook = download(client)

    assert [r[0] for r in workbook["Comenzi"].iter_rows(min_row=2, values_only=True)] == [
        visible.number
    ]


def test_export_is_recorded_in_audit_log(client, accounting_user):
    sent_order()
    client.force_login(accounting_user)

    download(client, q="CMD")

    event = BusinessEvent.objects.get(event_type=BusinessEvent.Type.ORDERS_EXPORTED)
    assert event.actor == accounting_user
    assert event.data == {"filtre": "căutare „CMD”"}


def test_export_button_keeps_filters(client, procurement_user):
    sent_order()
    client.force_login(procurement_user)

    html = client.get(
        reverse("purchasing:order_list"), {"status": "SENT", "page": 1}
    ).content.decode()

    assert f'href="{EXPORT}?status=SENT"' in html  # filtrele da, pagina nu
