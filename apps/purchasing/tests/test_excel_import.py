"""Importul comenzilor din Excel: conversii tolerante, validare pe rând, flux complet."""

import zipfile
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.audit.models import BusinessEvent
from apps.catalog.tests.factories import ProductFactory, SupplierFactory
from apps.purchasing import excel_import
from apps.purchasing.excel_import import (
    COLUMNS,
    ImportFileError,
    parse_date,
    parse_decimal,
    parse_vat_rate,
    plan_import,
    read_rows,
)
from apps.purchasing.models import PurchaseOrder

pytestmark = pytest.mark.django_db

TODAY = timezone.localdate()
SOON = TODAY + timedelta(days=10)
IMPORT = reverse("purchasing:order_import")
CONFIRM = reverse("purchasing:order_import_confirm")


# --- Utilitare de test ----------------------------------------------------------------------------


def xlsx(rows, header=None, sheet_title="Comenzi"):
    """Construiește în memorie un fișier .xlsx, ca un utilizator care completează șablonul."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append(header or [column.title for column in COLUMNS])
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile("comenzi.xlsx", buffer.getvalue())


def row(reference, supplier, product, quantity=10, price=100, expected=SOON, **extra):
    """Un rând în ordinea coloanelor din șablon."""
    values = {
        "reference": reference,
        "supplier": supplier.tax_id,
        "product": product.code,
        "quantity": quantity,
        "unit_price": price,
        "expected_date": expected,
        **extra,
    }
    return [values.get(column.key) for column in COLUMNS]


@pytest.fixture
def supplier():
    return SupplierFactory(tax_id="14399840", default_currency="RON")


@pytest.fixture
def products():
    return ProductFactory(code="MOT-65"), ProductFactory(code="RUL-6204", vat_rate=Decimal("11.00"))


# --- Conversii tolerante --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (12, Decimal("12")),
        (12.3, Decimal("12.3")),
        (0.30000000000000004, Decimal("0.3")),  # zgomotul de calcul al Excel-ului
        ("12,5", Decimal("12.5")),
        ("1.234,56", Decimal("1234.56")),  # stil românesc
        ("1,234.56", Decimal("1234.56")),  # stil englezesc
        ("1.500", Decimal("1500")),  # punct = separator de mii
        ("1,234,567", Decimal("1234567")),
        ("1 234,5", Decimal("1234.5")),
        (None, None),
    ],
)
def test_parse_decimal(value, expected):
    assert parse_decimal(value) == expected


@pytest.mark.parametrize("value", ["abc", "12,3,4.5x", 1e30])
def test_parse_decimal_rejects_garbage(value):
    with pytest.raises(ValueError):
        parse_decimal(value)


@pytest.mark.parametrize(
    "value", [date(2026, 10, 21), "2026-10-21", "21.10.2026", "21/10/2026", "21-10-2026", 46316]
)
def test_parse_date(value):
    assert parse_date(value) == date(2026, 10, 21)


def test_parse_date_rejects_garbage():
    with pytest.raises(ValueError, match="zz.ll.aaaa"):
        parse_date("31.02.2026")


@pytest.mark.parametrize(
    ("value", "expected"), [(21, "21.00"), ("21%", "21.00"), (0.11, "11.00"), (0, "0.00")]
)
def test_parse_vat_rate(value, expected):
    assert parse_vat_rate(value) == Decimal(expected)


# --- Citirea fișierului ---------------------------------------------------------------------------


def test_header_is_recognised_in_any_order_and_spelling(supplier, products):
    header = [
        "COD PRODUS",
        "cod fiscal furnizor",
        "Referinta comanda",
        "Cantitate",
        "Pret unitar",
        "Data estimata",
    ]
    upload = xlsx([[products[0].code, supplier.tax_id, "A-1", 3, 50, SOON]], header=header)

    rows = read_rows(upload)

    assert rows == [
        {
            "row": 2,
            "values": {
                "product": "MOT-65",
                "supplier": "14399840",
                "reference": "A-1",
                "quantity": 3,
                "unit_price": 50,
                "expected_date": SOON.isoformat(),
            },
        }
    ]


def test_missing_required_columns_are_named():
    with pytest.raises(ImportFileError, match="Cod produs, Cantitate"):
        read_rows(
            xlsx(
                [["A-1"]],
                header=["Referință comandă", "Cod fiscal furnizor", "Preț unitar", "Data estimată"],
            )
        )


def test_empty_rows_are_skipped_and_row_numbers_kept(supplier, products):
    upload = xlsx(
        [
            row("A-1", supplier, products[0]),
            [None] * len(COLUMNS),
            row("A-1", supplier, products[1]),
        ]
    )

    assert [r["row"] for r in read_rows(upload)] == [2, 4]  # numerele de rând din Excel


def test_file_without_lines_is_rejected():
    with pytest.raises(ImportFileError, match="nicio linie"):
        read_rows(xlsx([]))


def test_not_an_excel_file_is_rejected():
    with pytest.raises(ImportFileError, match="nu este un Excel"):
        read_rows(SimpleUploadedFile("comenzi.xlsx", b"nu sunt un fisier excel"))


def test_too_many_rows(monkeypatch, supplier, products):
    monkeypatch.setattr(excel_import, "MAX_ROWS", 2)

    with pytest.raises(ImportFileError, match="mai mult de 2"):
        read_rows(xlsx([row("A", supplier, products[0])] * 3))


def test_zip_bomb_is_rejected(monkeypatch):
    """Un .xlsx e o arhivă: refuzăm fișierele care s-ar „umfla” enorm la dezarhivare."""
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", "0" * 2_000_000)  # 2 MB comprimat în câțiva KB
    monkeypatch.setattr(excel_import, "MAX_UNCOMPRESSED_BYTES", 1_000_000)

    with pytest.raises(ImportFileError, match="prea mare"):
        read_rows(SimpleUploadedFile("bomba.xlsx", buffer.getvalue()))


# --- Validare și grupare --------------------------------------------------------------------------


def plan_for(user, rows):
    return plan_import(read_rows(xlsx(rows)), user)


def test_rows_with_same_reference_form_one_order(procurement_user, supplier, products):
    later = SOON + timedelta(days=5)
    plan = plan_for(
        procurement_user,
        [
            row("A-1", supplier, products[0], quantity=2, price=1500),
            row("A-1", supplier, products[1], quantity="10", price="25,5", expected=later),
            row("B-7", supplier, products[0]),
        ],
    )

    assert plan.is_valid, plan.error_messages()
    first, second = plan.orders
    assert first.reference == "A-1" and len(first.lines) == 2
    assert first.order.currency == "RON"  # moneda implicită a furnizorului
    assert first.order.expected_date == SOON  # cel mai apropiat termen al liniilor
    assert first.lines[1][1].vat_rate == Decimal("11.00")  # cota implicită a produsului
    assert first.net_amount == Decimal("3255.00")
    assert second.reference == "B-7"


@pytest.mark.parametrize(
    ("change", "message"),
    [
        ({"product": "NU-EXISTA"}, "produsul „NU-EXISTA” nu există"),
        ({"supplier": "99999999"}, "furnizorul cu codul fiscal „99999999” nu există"),
        ({"quantity": 0}, "Cantitate: Cantitatea trebuie să fie pozitivă"),
        ({"quantity": "doi"}, "Cantitate: „doi” nu este un număr"),
        ({"unit_price": -5}, "Preț unitar: Prețul nu poate fi negativ"),
        ({"currency": "USD"}, "„USD” nu este acceptată"),
        ({"vat_rate": 19}, "Cotă TVA (%)"),
        ({"expected_date": None}, "Data estimată: lipsește"),
        ({"expected_date": TODAY - timedelta(days=1)}, "nu poate fi înaintea datei comenzii"),
        ({"reference": "R" * 51}, "Referință comandă"),
    ],
)
def test_row_errors_are_reported_with_row_number(
    procurement_user, supplier, products, change, message
):
    values = dict(zip([c.key for c in COLUMNS], row("A-1", supplier, products[0]), strict=True))
    values.update(change)

    plan = plan_for(procurement_user, [[values[c.key] for c in COLUMNS]])

    assert not plan.is_valid
    assert any(e.startswith("Rândul 2:") and message in e for e in plan.error_messages()), (
        plan.error_messages()
    )


def test_inactive_product_and_supplier_are_rejected(procurement_user, supplier, products):
    products[0].is_active = False
    products[0].save()

    plan = plan_for(procurement_user, [row("A-1", supplier, products[0])])

    assert "inactiv" in plan.error_messages()[0]


def test_one_order_cannot_mix_suppliers(procurement_user, supplier, products):
    other = SupplierFactory()

    plan = plan_for(
        procurement_user, [row("A-1", supplier, products[0]), row("A-1", other, products[1])]
    )

    assert not plan.is_valid
    assert "Rândul 3: Cod fiscal furnizor diferă de rândul 2" in plan.error_messages()[0]


def test_all_errors_are_reported_at_once(procurement_user, supplier, products):
    """Utilizatorul corectează tot fișierul dintr-o dată, nu eroare cu eroare."""
    plan = plan_for(
        procurement_user,
        [
            row("A-1", supplier, products[0], quantity=0),
            row("B-1", supplier, products[0], price="abc"),
            row("C-1", supplier, products[0]),  # valid
        ],
    )

    rows_with_errors = {message.split(":")[0] for message in plan.error_messages()}
    assert rows_with_errors == {"Rândul 2", "Rândul 3"}
    assert not plan.is_valid  # un fișier cu erori nu se importă deloc, nici parțial


# --- Fluxul complet prin interfață ----------------------------------------------------------------


def test_full_import_flow(client, procurement_user, supplier, products):
    client.force_login(procurement_user)

    preview = client.post(
        IMPORT,
        {"file": xlsx([row("A-1", supplier, products[0]), row("A-1", supplier, products[1])])},
    )
    assert "Se vor crea" in preview.content.decode()
    assert not PurchaseOrder.objects.exists()  # previzualizarea nu salvează nimic

    response = client.post(CONFIRM, {"token": preview.context["token"]}, follow=True)

    order = PurchaseOrder.objects.get()
    assert order.status == PurchaseOrder.Status.DRAFT
    assert order.owner == procurement_user
    assert order.external_reference == "A-1"
    assert order.lines.count() == 2
    assert "Au fost importate 1 comenzi" in response.content.decode()
    event = BusinessEvent.objects.get(event_type=BusinessEvent.Type.ORDERS_IMPORTED)
    assert event.actor == procurement_user
    assert event.data["comenzi"] == order.number


def test_confirmation_works_only_once(client, procurement_user, supplier, products):
    client.force_login(procurement_user)
    preview = client.post(IMPORT, {"file": xlsx([row("A-1", supplier, products[0])])})
    token = preview.context["token"]

    client.post(CONFIRM, {"token": token})
    response = client.post(CONFIRM, {"token": token}, follow=True)

    assert PurchaseOrder.objects.count() == 1
    assert "a fost deja efectuat" in response.content.decode()


def test_same_file_cannot_be_imported_twice(client, procurement_user, supplier, products):
    client.force_login(procurement_user)
    upload = [row("A-1", supplier, products[0])]
    first = client.post(IMPORT, {"file": xlsx(upload)})
    client.post(CONFIRM, {"token": first.context["token"]})

    second = client.post(IMPORT, {"file": xlsx(upload)})

    assert "Există deja o comandă cu această referință externă" in second.content.decode()
    assert PurchaseOrder.objects.count() == 1


def test_data_changed_between_preview_and_confirmation(
    client, procurement_user, supplier, products
):
    client.force_login(procurement_user)
    preview = client.post(IMPORT, {"file": xlsx([row("A-1", supplier, products[0])])})
    products[0].is_active = False  # cineva dezactivează produsul între timp
    products[0].save()

    response = client.post(CONFIRM, {"token": preview.context["token"]})

    assert not PurchaseOrder.objects.exists()
    assert "Datele s-au schimbat" in response.content.decode()


def test_upload_rejects_other_file_types(client, procurement_user):
    client.force_login(procurement_user)

    response = client.post(IMPORT, {"file": SimpleUploadedFile("comenzi.csv", b"a;b")})

    assert "doar fișiere Excel .xlsx" in response.content.decode()


def test_upload_shows_file_level_errors(client, procurement_user):
    client.force_login(procurement_user)

    response = client.post(IMPORT, {"file": xlsx([["A-1"]], header=["Referință comandă"])})

    assert "Lipsesc coloanele" in response.content.decode()


# --- Șablonul -------------------------------------------------------------------------------------


def test_template_contains_current_products_and_suppliers(
    client, procurement_user, supplier, products
):
    client.force_login(procurement_user)

    response = client.get(reverse("purchasing:order_import_template"))

    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Comenzi", "Instrucțiuni", "Furnizori", "Produse"]
    assert [c.value for c in workbook["Comenzi"][1]] == [c.title for c in COLUMNS]
    assert {r[0] for r in workbook["Produse"].iter_rows(min_row=2, values_only=True)} == {
        "MOT-65",
        "RUL-6204",
    }
    assert workbook["Comenzi"].data_validations.dataValidation  # listele derulante


def test_template_round_trip(client, procurement_user, supplier, products):
    """Șablonul descărcat, completat și încărcat înapoi trebuie să funcționeze."""
    client.force_login(procurement_user)
    workbook = load_workbook(
        BytesIO(client.get(reverse("purchasing:order_import_template")).content)
    )
    workbook["Comenzi"].append(row("T-1", supplier, products[0]))
    buffer = BytesIO()
    workbook.save(buffer)

    preview = client.post(IMPORT, {"file": SimpleUploadedFile("sablon.xlsx", buffer.getvalue())})

    assert preview.context["plan"].is_valid, preview.context["plan"].error_messages()


def test_errors_in_one_row_do_not_hide_errors_in_other_rows_of_the_same_order(
    procurement_user, supplier, products
):
    """Rândul 2 are un produs inexistent; rândul 3 (aceeași comandă) are cantitate 0 și alt
    furnizor. Toate trei problemele trebuie raportate de la prima încărcare."""
    other = SupplierFactory()
    unknown_product = row("A-1", supplier, products[0])
    unknown_product[[c.key for c in COLUMNS].index("product")] = "NU-EXISTA"
    plan = plan_for(procurement_user, [unknown_product, row("A-1", other, products[1], quantity=0)])

    messages = plan.error_messages()
    assert any(m.startswith("Rândul 2: produsul „NU-EXISTA”") for m in messages)
    assert any(m.startswith("Rândul 3: Cantitate:") for m in messages)
    assert any(m.startswith("Rândul 3: Cod fiscal furnizor diferă de rândul 2") for m in messages)
