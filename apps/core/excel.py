"""Piese comune pentru fișierele Excel generate de aplicație (exporturi, șabloane).

Valorile se scriu ca numere și date reale, nu ca text: în Excel se pot face
sume, filtre și sortări. Formatele de mai jos sunt independente de limbă:
un Excel în română afișează automat 1.234,56 și 21.10.2026.
"""

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Formate de afișare
MONEY = "#,##0.00"
PRICE = "#,##0.00##"  # minim 2, maxim 4 zecimale
QUANTITY = "General"  # 12 sau 12,5 – fără virgulă inutilă la numere întregi
DATE = "DD.MM.YYYY"
DATETIME = "DD.MM.YYYY HH:MM"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")  # culoarea aplicației


@dataclass(frozen=True)
class Column:
    title: str
    number_format: str | None = None
    width: int | None = None  # None = calculată din conținut


def excel_value(value):
    """Valori pe care Excel le înțelege: fără fus orar, fără tipuri Django."""
    if isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None) if timezone.is_aware(value) else value
    if value is None or isinstance(value, date | int | float | Decimal):
        return value  # numerele (inclusiv Decimal) și datele rămân ca atare
    return str(value)


def style_header(sheet, row=1):
    for cell in sheet[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 30


def write_table(sheet, columns, rows):
    """Scrie un tabel: antet formatat, valori, formate de număr, lățimi, filtre."""
    sheet.append([column.title for column in columns])
    style_header(sheet)
    widths = [len(column.title) for column in columns]

    for row in rows:
        sheet.append([excel_value(value) for value in row])
        for index, value in enumerate(row):
            text = value.strftime("%d.%m.%Y") if isinstance(value, date) else str(value or "")
            widths[index] = max(widths[index], len(text))

    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        auto_width = min(max(widths[index - 1], 8) + 2, 60)  # între 8 și 60 de caractere
        sheet.column_dimensions[letter].width = column.width or auto_width
        if column.number_format:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = column.number_format

    sheet.freeze_panes = "A2"  # antetul rămâne vizibil la derulare
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions  # săgețile de filtrare din Excel


def xlsx_response(workbook, filename):
    """Răspuns HTTP care descarcă registrul ca fișier .xlsx."""
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
